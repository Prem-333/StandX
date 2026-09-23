"""Actual FastAPI integration against the Phase 2 seed and persistent local SQL."""
from io import BytesIO
import json
from pathlib import Path
import socket
import sys
from uuid import uuid4
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
sys.stdout.reconfigure(encoding='utf-8')

original_connect=socket.socket.connect
attempts=[]
def local_only(sock,address):
    if isinstance(address,tuple) and address[0] not in ('127.0.0.1','localhost','::1'):
        attempts.append(str(address));raise RuntimeError('External network disabled in integration demo')
    return original_connect(sock,address)
socket.socket.connect=local_only

from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject,DictionaryObject,NameObject
from services.api.app import create_app
from services.api.runtime import create_runtime
from services.nlp.multilingual import QueryNormalizer,configuration
from services.ingestion.records import ROOT
from phase7_demo import evaluate

KEY='local-integration-'+uuid4().hex
headers={'X-API-Key':KEY}
runtime=create_runtime()
checks=[]

def check(name,response,status=200):
    assert response.status_code==status,(name,response.status_code,response.text[:1200])
    checks.append({'check':name,'status':response.status_code})
    return response.json()

try:
    with TestClient(create_app(runtime,{'test-officer':KEY,'other-officer':KEY+'-other'},rate_limit=200)) as client:
        check('unauthorized',client.get('/v1/health'),401)
        check('invalid API key',client.get('/v1/health',headers={'X-API-Key':'wrong-key'}),401)
        check('health',client.get('/v1/health',headers=headers))
        check('standard detail',client.get('/v1/standards/IS%209550:2024',headers=headers))
        check('allied standards',client.get('/v1/standards/IS%209550:2024/allied',headers=headers))
        check('unknown standard',client.get('/v1/standards/IS%20999999:2099',headers=headers),404)
        check('invalid hops',client.get('/v1/standards/IS%209550:2024/allied?max_hops=9',headers=headers),422)
        for category in ('bright_steel_bars','laptop_notebook_tablet','silver_jewellery_artefacts','unknown_category'):
            check('certifications '+category,client.get('/v1/certifications/'+category,headers=headers))
        recommendation=check('text recommendation',client.post('/v1/recommend',headers=headers,json={
            'text':'IS 9550:2024','product_category':'bright_steel_bars',
            'certification_context':{'domestic_supply':True,'exemption_claimed':False}}))
        assert recommendation['primary_standards'][0]['record_id']=='bis-10'
        assert recommendation['certification_requirements'][0]['requirement']=='mandatory'
        assert recommendation['primary_standards'][0]['evidence'][0]['is_number']=='IS 9550:2024'
        partial=check('multiple literal references preserve unresolved part',client.post('/v1/recommend',headers=headers,json={
            'text':'IS 7524 optical tests and IS 5983 specification'}))
        assert partial['status']=='partial_review_required' and partial['unresolved_identifiers']==['IS 7524']
        assert partial['primary_standards'][0]['is_number']=='IS 5983:1980'
        history=check('owned history',client.get('/v1/history',headers=headers))
        assert any(r['recommendation_id']==recommendation['recommendation_id'] for r in history['items'])
        saved=check('restore saved report',client.get('/v1/history/'+recommendation['recommendation_id'],headers=headers))
        assert saved==recommendation
        other={'X-API-Key':KEY+'-other'}
        assert check('other officer isolated history',client.get('/v1/history',headers=other))['items']==[]
        check('other officer cannot read report',client.get('/v1/history/'+recommendation['recommendation_id'],headers=other),404)
        check('other officer cannot write feedback',client.post('/v1/feedback',headers=other,json={
            'recommendation_id':recommendation['recommendation_id'],'decision':'confirm','record_id':'bis-10'}),404)
        directory=check('directory filtered',client.get('/v1/standards?q=9550',headers=headers))
        assert directory['total']==1 and directory['items'][0]['evidence'][0]['record_id']=='bis-10'
        check('directory invalid page',client.get('/v1/standards?limit=10000',headers=headers),422)
        system=check('active configuration',client.get('/v1/system',headers=headers))
        assert system['verified_records']==20 and 'embedding_model' in system['models']
        for bad in ('   ','bad\x00text','bad\ud800text'):
            check('invalid Unicode or empty text',client.post('/v1/recommend',headers={**headers,'Content-Type':'application/json'},content=json.dumps({'text':bad})),422)
        check('chunked oversized feedback',client.post('/v1/feedback',headers={**headers,'Content-Type':'application/json'},
            content=iter([b'x'*9000,b'x'*9000])),413)
        check('feedback',client.post('/v1/feedback',headers=headers,json={
            'recommendation_id':recommendation['recommendation_id'],'decision':'confirm','record_id':'bis-10','comment':'Integration test confirmation'}),201)
        check('feedback missing recommendation',client.post('/v1/feedback',headers=headers,json={
            'recommendation_id':str(uuid4()),'decision':'reject'}),404)
        check('reject unrelated record forbidden',client.post('/v1/feedback',headers=headers,json={
            'recommendation_id':recommendation['recommendation_id'],'decision':'reject','record_id':'bis-04'}),422)
        check('feedback unknown KB record',client.post('/v1/feedback',headers=headers,json={
            'recommendation_id':recommendation['recommendation_id'],'decision':'correct','record_id':'invented'}),422)
        check('invalid input',client.post('/v1/recommend',headers=headers,json={'text':'','top_k':50}),422)
        doc=Document();doc.add_paragraph('Supply wooden bedside tables.');doc.add_paragraph('Provide cloud payroll software subscription.')
        data=BytesIO();doc.save(data)
        result=check('DOCX tender upload',client.post('/v1/recommend',headers=headers,
            files={'file':('tender.docx',data.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}))
        assert result['kind']=='tender' and result['phrase_results']
        writer=PdfWriter();page=writer.add_blank_page(width=500,height=200)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 100 Td (Supply wooden bedside tables.) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(stream)
        pdf=BytesIO();writer.write(pdf)
        check('PDF tender upload',client.post('/v1/recommend',headers=headers,files={'file':('tender.pdf',pdf.getvalue(),'application/pdf')}))
        check('bad upload',client.post('/v1/recommend',headers=headers,files={'file':('tender.exe',b'no','application/octet-stream')}),422)
        check('oversized body',client.post('/v1/recommend',headers={**headers,'Content-Type':'application/json'},content=b'x'*(5*1024*1024+65537)),413)
        hindi=check('Hindi API boundary',client.post('/v1/recommend',headers=headers,json={'text':'लकड़ी की बेडसाइड मेज'}))
        assert hindi['normalization']['status']=='translated'
        assert hindi['primary_standards'][0]['is_number']=='IS 6188:1988'
        original_normalizer=runtime.normalizer
        runtime.normalizer=QueryNormalizer({**configuration(),'enabled':False})
        fallback=check('translation absent fallback',client.post('/v1/recommend',headers=headers,json={'text':'लकड़ी की मेज','language_hint':'hi'}))
        assert not fallback['primary_standards'] and fallback['normalization']['status']=='english_only_fallback'
        runtime.normalizer=original_normalizer
        schema=check('OpenAPI',client.get('/openapi.json'))
        docs=client.get('/docs')
        assert docs.status_code==200 and 'cdn.jsdelivr.net' not in docs.text
        assert '/static/swagger-ui-bundle.js' in docs.text
        import re
        nonce=re.search(r'<script nonce="([^"]+)"',docs.text).group(1)
        assert f"'nonce-{nonce}'" in docs.headers['content-security-policy']
        assert nonce not in client.get('/docs').text
        asset=client.get('/static/swagger-ui-bundle.js')
        assert asset.status_code==200
        checks.append({'check':'offline Swagger UI and assets','status':200})
        assert set(schema['paths']['/v1/recommend']['post']['requestBody']['content'])=={'application/json','multipart/form-data'}
        assert 'APIKeyHeader' in schema['components']['securitySchemes']
        (ROOT/'docs/openapi.json').write_text(json.dumps(schema,indent=2),encoding='utf-8')
    with TestClient(create_app(runtime,{'test-officer':KEY},rate_limit=1)) as client:
        check('rate first request',client.get('/v1/health',headers=headers))
        response=client.get('/v1/health',headers=headers);check('rate limited',response,429)
        assert response.headers['Retry-After']
    multilingual=evaluate(runtime)
    assert multilingual['translation_completed']==8,'One or more translations failed; inspect report'
    assert multilingual['top5']==8,'A multilingual smoke query missed the expected record; inspect report'
    assert not attempts
    report={'checks':checks,'passed':len(checks),'external_connection_attempts':attempts,
        'multilingual_top1':multilingual['top1'],'multilingual_top5':multilingual['top5'],
        'audit_rows':runtime.connection.execute('SELECT count(*) FROM kb.recommendations_log').fetchone()[0],
        'feedback_rows':runtime.connection.execute('SELECT count(*) FROM kb.user_feedback').fetchone()[0],
        'environment':'FastAPI TestClient, actual cached models, persistent local Qdrant and PGlite PostgreSQL; Docker not executed'}
    (ROOT/'data/processed/api_integration_report.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report,indent=2),flush=True)
finally:runtime.close()
